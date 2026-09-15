import io
import re
from typing import Optional, Set
from fastapi import FastAPI, File, UploadFile, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Presentation Parser Imports
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

app = FastAPI(title="Document & Presentation Compliance Auditor")

# Enable CORS for local development and deployed frontend origins
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

WORD_ALIGN_MAP = {
    "LEFT": WD_ALIGN_PARAGRAPH.LEFT,
    "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT,
    "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "CENTER": WD_ALIGN_PARAGRAPH.CENTER
}

PPT_ALIGN_MAP = {
    "LEFT": PP_ALIGN.LEFT,
    "RIGHT": PP_ALIGN.RIGHT,
    "JUSTIFY": PP_ALIGN.JUSTIFY,
    "CENTER": PP_ALIGN.CENTER
}

KNOWN_HEADINGS = {
    "abstract", "acknowledgements", "acknowledgment", "table of contents", 
    "contents", "introduction", "background", "methodology", "results", 
    "discussion", "conclusion", "references", "appendix", "executive summary",
    "project description", "problem statement", "literature review", "system design",
    "implementation", "testing and evaluation", "future work"
}

BODY_REFERENCE_VERBS = (
    "represents", "shows", "illustrates", "displays", "lists", "depicts",
    "gives", "demonstrates", "highlights", "summarizes", "compares", "is",
    "was", "are", "were", "contains", "provides", "describes", "details",
    "indicates", "outlines", "presents"
)

def is_true(val: Optional[str]) -> bool:
    if val is None:
        return False
    return str(val).lower() in ("true", "1", "yes")

def normalize_font_name(font_name: Optional[str]) -> str:
    if not font_name:
        return ""
    return re.sub(r'\s*\([^)]*\)', '', font_name).strip()

def parse_excluded_pages(pages_str: Optional[str]) -> Set[int]:
    """Parses inputs like '1, 2, 4-6' into a set of integers {1, 2, 4, 5, 6}."""
    if not pages_str:
        return set()
    
    excluded = set()
    parts = [p.strip() for p in pages_str.split(",") if p.strip()]
    for part in parts:
        if "-" in part:
            try:
                start, end = map(int, part.split("-"))
                excluded.update(range(start, end + 1))
            except ValueError:
                continue
        else:
            try:
                excluded.add(int(part))
            except ValueError:
                continue
    return excluded

def get_column_count(section) -> int:
    cols_elements = section._sectPr.xpath('./w:cols')
    if cols_elements:
        num = cols_elements[0].get(qn('w:num'))
        if num:
            return int(num)
    return 1


def verify_presentation_deck(
    contents: bytes,
    preset: str,
    check_reading_order: bool,
    check_boundary_overflow: bool,
    check_title_alignment: bool,
    check_text_justification: bool,
    target_alignment: str,
    title_slide_font: Optional[str],
    title_slide_font_size: Optional[float],
    body_heading_font: Optional[str],
    body_heading_font_size: Optional[float],
    body_slide_font: Optional[str],
    body_slide_font_size: Optional[float],
    aspect_ratio: Optional[str],
    excluded_pages: Set[int] = set()
):
    prs = Presentation(io.BytesIO(contents))
    errors_set = set()
    passed = []
    slide_analysis = []

    slide_width_in = prs.slide_width.inches
    slide_height_in = prs.slide_height.inches

    if preset == "custom_ppt":
        t_title_font = title_slide_font or "Calibri"
        t_title_size = title_slide_font_size or 44.0
        t_heading_font = body_heading_font or "Times New Roman"
        t_heading_size = body_heading_font_size or 24.0
        t_body_font = body_slide_font or "Times New Roman"
        t_body_size = body_slide_font_size or 20.0
        preset_label = "Custom Slide Ruleset"
    elif preset == "academic_slide":
        t_title_font = title_slide_font or "Times New Roman"
        t_title_size = title_slide_font_size or 36.0
        t_heading_font = body_heading_font or "Arial"
        t_heading_size = body_heading_font_size or 24.0
        t_body_font = body_slide_font or "Arial"
        t_body_size = body_slide_font_size or 18.0
        preset_label = "Academic Defense Standard"
    elif preset == "deck_4_3":
        t_title_font = title_slide_font or "Arial"
        t_title_size = title_slide_font_size or 32.0
        t_heading_font = body_heading_font or "Arial"
        t_heading_size = body_heading_font_size or 22.0
        t_body_font = body_slide_font or "Arial"
        t_body_size = body_slide_font_size or 16.0
        preset_label = "Standard 4:3 Deck"
    else: 
        t_title_font = title_slide_font or "Arial"
        t_title_size = title_slide_font_size or 36.0
        t_heading_font = body_heading_font or "Arial"
        t_heading_size = body_heading_font_size or 24.0
        t_body_font = body_slide_font or "Arial"
        t_body_size = body_slide_font_size or 18.0
        preset_label = "Widescreen 16:9 Deck"

    target_align_enum = PPT_ALIGN_MAP.get(target_alignment.upper(), PP_ALIGN.LEFT)

    title_font_clean = True
    heading_font_clean = True
    body_font_clean = True
    justification_clean = True
    boundary_clean = True
    reading_order_clean = True

    for idx, slide in enumerate(prs.slides, start=1):
        if idx in excluded_pages:
            continue

        layout_name = slide.slide_layout.name.lower().strip()
        is_cover_title_slide = (idx == 1) or (layout_name in ["title slide", "section header"])
        shapes_in_slide = []
        sections = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            shape_top = shape.top.inches
            shape_left = shape.left.inches
            shape_width = shape.width.inches
            shape_height = shape.height.inches
            
            is_title_shape = ("title" in shape.name.lower() or "header" in shape.name.lower()) and (shape_top < 2.5)

            shapes_in_slide.append({
                "name": shape.name,
                "top": shape_top,
                "left": shape_left,
                "width": shape_width,
                "height": shape_height
            })

            if check_boundary_overflow:
                if (shape_left + shape_width > slide_width_in + 0.05) or (shape_top + shape_height > slide_height_in + 0.05):
                    errors_set.add(f"Slide {idx}: Content box '{shape.name}' exceeds slide boundary canvas.")
                    boundary_clean = False

            for p in shape.text_frame.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                text_lower = text.lower()
                is_caption = text_lower.startswith(("figure", "fig.", "fig ", "table", "source:", "available:", "doi:", "http://", "https://", "["))

                if is_title_shape and not is_caption:
                    for run in p.runs:
                        font_name = normalize_font_name(run.font.name)
                        if font_name and font_name.lower() != t_title_font.lower():
                            errors_set.add(f"Slide {idx} (Slide Title '{text[:25]}'): Font '{font_name}' does not match target title font '{t_title_font}'.")
                            title_font_clean = False
                            break
                        if run.font.size and run.font.size.pt != t_title_size:
                            errors_set.add(f"Slide {idx} (Slide Title '{text[:25]}'): Font size {run.font.size.pt}pt does not match target title size {t_title_size}pt.")
                            title_font_clean = False
                            break

                elif not is_caption and p.level == 0 and (any(r.font.bold is True for r in p.runs) or p.font.bold is True or (len(text) <= 55 and not text.endswith(('.', ':', ';', ',')))):
                    for run in p.runs:
                        font_name = normalize_font_name(run.font.name)
                        if font_name and font_name.lower() != t_heading_font.lower():
                            errors_set.add(f"Slide {idx} (Sub-Heading '{text[:25]}...'): Font '{font_name}' does not match target heading font '{t_heading_font}'.")
                            heading_font_clean = False
                            break
                        if run.font.size and run.font.size.pt != t_heading_size:
                            errors_set.add(f"Slide {idx} (Sub-Heading '{text[:25]}...'): Font size {run.font.size.pt}pt does not match target heading size {t_heading_size}pt.")
                            heading_font_clean = False
                            break

                else:
                    for run in p.runs:
                        font_name = normalize_font_name(run.font.name)
                        if font_name and font_name.lower() != t_body_font.lower():
                            errors_set.add(f"Slide {idx} (Body Text): Font '{font_name}' does not match target '{t_body_font}'.")
                            body_font_clean = False
                            break
                        if run.font.size and run.font.size.pt != t_body_size:
                            errors_set.add(f"Slide {idx} (Body Text): Size {run.font.size.pt}pt does not match target {t_body_size}pt.")
                            body_font_clean = False
                            break

                    if check_text_justification:
                        if p.alignment is not None and p.alignment != target_align_enum:
                            errors_set.add(f"Slide {idx}: Text paragraph starting with '{text[:25]}...' is not using required {target_alignment} alignment.")
                            justification_clean = False

            sections.append({
                "name": shape.name,
                "top_in": round(shape_top, 2),
                "width_in": round(shape_width, 2)
            })

        in_order = True
        if check_reading_order and shapes_in_slide:
            sorted_shapes = sorted(shapes_in_slide, key=lambda s: s["top"])
            top_shape_name = sorted_shapes[0]["name"].lower()
            if "title" not in top_shape_name and "header" not in top_shape_name:
                in_order = False
                reading_order_clean = False
                errors_set.add(f"Slide {idx}: Reading order issue — top-most block '{sorted_shapes[0]['name']}' is not a Header/Title.")

        slide_analysis.append({
            "slide_number": idx,
            "is_title_slide": is_cover_title_slide,
            "in_order": in_order,
            "sections": sections
        })

    def parse_slide_num(err_msg: str) -> tuple:
        match = re.search(r'Slide (\d+)', err_msg)
        slide_num = int(match.group(1)) if match else 9999
        return (slide_num, err_msg)

    errors = sorted(list(errors_set), key=parse_slide_num)

    if title_font_clean and len(prs.slides) > 0:
        passed.append(f"Title slide typography matches required standard ({t_title_font}, {t_title_size}pt).")
    if heading_font_clean and len(prs.slides) > 0:
        passed.append(f"Body sub-headings typography matches target standard ({t_heading_font}, {t_heading_size}pt).")
    if body_font_clean and len(prs.slides) > 0:
        passed.append(f"Regular body text matches target standard ({t_body_font}, {t_body_size}pt).")
    if check_text_justification and justification_clean and len(prs.slides) > 0:
        passed.append(f"Body paragraph justification complies with required alignment ({target_alignment}).")
    if check_boundary_overflow and boundary_clean and len(prs.slides) > 0:
        passed.append("All slide content frames remain inside physical slide boundaries.")
    if check_reading_order and reading_order_clean and len(prs.slides) > 0:
        passed.append("Vertical reading order across slides correctly places header elements top-most.")

    total_checks = len(errors) + len(passed)
    score = 100 if total_checks == 0 else int((len(passed) / total_checks) * 100)

    return {
        "preset_name": preset_label,
        "compliance_score": score,
        "errors": errors,
        "passed": passed,
        "slide_analysis": slide_analysis
    }


# Dual decorators to catch requests with or without /api prefix
@app.post("/api/verify")
@app.post("/verify")
async def verify_document(
    file: UploadFile = File(...),
    preset: str = Form(...),
    file_type: Optional[str] = Form("document"),
    
    # Custom Document Settings
    custom_title_font: Optional[str] = Form("Times New Roman"),
    custom_title_size: Optional[float] = Form(24.0),
    custom_subtitle_font: Optional[str] = Form("Times New Roman"),
    custom_subtitle_size: Optional[float] = Form(14.0),
    custom_subtitle_bold: Optional[str] = Form("true"),
    
    # Level 3 Heading Options
    custom_subsub_font: Optional[str] = Form("Times New Roman"),
    custom_subsub_size: Optional[float] = Form(12.0),
    custom_subsub_bold: Optional[str] = Form("true"),
    custom_subsub_italic: Optional[str] = Form("true"),

    custom_font: Optional[str] = Form("Times New Roman"),
    custom_size: Optional[float] = Form(12.0),
    custom_margin_top: Optional[float] = Form(1.0),
    custom_margin_bottom: Optional[float] = Form(1.0),
    custom_margin_left: Optional[float] = Form(1.0),
    custom_margin_right: Optional[float] = Form(1.0),
    custom_line_spacing: Optional[float] = Form(1.15),
    
    # Presentation Settings
    check_reading_order: Optional[str] = Form("false"),
    check_boundary_overflow: Optional[str] = Form("false"),
    check_title_alignment: Optional[str] = Form("false"),
    check_text_justification: Optional[str] = Form("false"),
    target_alignment: Optional[str] = Form("LEFT"),
    title_slide_font: Optional[str] = Form("Georgia"),
    title_slide_font_size: Optional[float] = Form(40.0),
    body_heading_font: Optional[str] = Form("Arial"),
    body_heading_font_size: Optional[float] = Form(24.0),
    body_slide_font: Optional[str] = Form("Arial"),
    body_slide_font_size: Optional[float] = Form(18.0),
    aspect_ratio: Optional[str] = Form("16:9"),
    excluded_pages: Optional[str] = Form(""),
):
    filename = file.filename.lower()
    contents = await file.read()
    excluded_set = parse_excluded_pages(excluded_pages)

    if filename.endswith(('.ppt', '.pptx')) or file_type == "presentation":
        try:
            return verify_presentation_deck(
                contents=contents, preset=preset,
                check_reading_order=is_true(check_reading_order),
                check_boundary_overflow=is_true(check_boundary_overflow),
                check_title_alignment=is_true(check_title_alignment),
                check_text_justification=is_true(check_text_justification),
                target_alignment=target_alignment or "LEFT",
                title_slide_font=title_slide_font, title_slide_font_size=title_slide_font_size,
                body_heading_font=body_heading_font, body_heading_font_size=body_heading_font_size,
                body_slide_font=body_slide_font, body_slide_font_size=body_slide_font_size,
                aspect_ratio=aspect_ratio,
                excluded_pages=excluded_set
            )
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse presentation file: {str(e)}")

    elif filename.endswith('.docx'):
        try:
            doc = Document(io.BytesIO(contents))
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse .docx file: {str(e)}")

        # Configure targets based on preset
        if preset == "custom":
            target_title_font = custom_title_font or "Times New Roman"
            target_title_size = custom_title_size or 24.0
            
            target_subtitle_font = custom_subtitle_font or "Times New Roman"
            target_subtitle_size = custom_subtitle_size or 14.0
            require_subtitle_bold = is_true(custom_subtitle_bold)
            
            target_subsub_font = custom_subsub_font or "Times New Roman"
            target_subsub_size = custom_subsub_size or 12.0
            require_subsub_bold = is_true(custom_subsub_bold)
            require_subsub_italic = is_true(custom_subsub_italic)

            target_font = custom_font or "Times New Roman"
            target_size = custom_size or 12.0
            target_margins = {"top": custom_margin_top or 1.0, "bottom": custom_margin_bottom or 1.0, "left": custom_margin_left or 1.0, "right": custom_margin_right or 1.0}
            target_cols = 1
            target_spacing = custom_line_spacing or 1.15
            preset_label = "Custom Ruleset"
            
        elif preset == "ieee":
            target_title_font = "Times New Roman"
            target_title_size = 24.0
            target_subtitle_font = "Times New Roman"
            target_subtitle_size = 10.0
            require_subtitle_bold = False
            
            target_subsub_font = "Times New Roman"
            target_subsub_size = 10.0
            require_subsub_bold = False
            require_subsub_italic = True

            target_font = "Times New Roman"
            target_size = 10.0
            target_margins = {"top": 0.75, "bottom": 0.75, "left": 0.625, "right": 0.625}
            target_cols = 2
            target_spacing = 1.0
            preset_label = "IEEE Standard"
            
        else:  # APA 7
            target_title_font = "Times New Roman"
            target_title_size = 12.0
            target_subtitle_font = "Times New Roman"
            target_subtitle_size = 12.0
            require_subtitle_bold = True
            
            target_subsub_font = "Times New Roman"
            target_subsub_size = 12.0
            require_subsub_bold = True
            require_subsub_italic = True

            target_font = "Times New Roman"
            target_size = 12.0
            target_margins = {"top": 1.0, "bottom": 1.0, "left": 1.0, "right": 1.0}
            target_cols = 1
            target_spacing = 2.0
            preset_label = "APA 7th Edition"

        errors = []
        passed = []

        current_page = 1
        paragraph_counter = 1
        valid_paragraphs = 0
        current_sec_idx = 0
        margin_checked = False
        
        title_clean = True
        subtitle_clean = True
        subsub_clean = True
        font_clean = True
        size_clean = True
        spacing_clean = True
        indent_clean = True
        justification_clean = True
        caption_clean = True
        has_captions = False
        has_subsubs = False

        page_height_inches = 11.0 
        if doc.sections and doc.sections[0].page_height:
            page_height_inches = doc.sections[0].page_height.inches
            
        top_m = target_margins["top"]
        bot_m = target_margins["bottom"]
        usable_height = page_height_inches - (top_m + bot_m)
        
        line_height_inches = (target_size / 72.0) * target_spacing
        max_lines_per_page = max(15, int(usable_height / line_height_inches))
        current_page_line_count = 0

        target_word_align_enum = WORD_ALIGN_MAP.get((target_alignment or "LEFT").upper(), WD_ALIGN_PARAGRAPH.LEFT)
        prev_was_chapter = False

        for p in doc.paragraphs:
            page_breaks = p._element.xpath('.//w:lastRenderedPageBreak | .//w:br[@w:type="page"]')
            has_section_break = bool(p._element.xpath('.//w:sectPr'))
            
            text = p.text.strip()
            if not text:
                if page_breaks:
                    current_page += len(page_breaks)
                    paragraph_counter = 1
                    current_page_line_count = 0
                if has_section_break:
                    current_sec_idx = min(current_sec_idx + 1, len(doc.sections) - 1)
                continue

            # Run section & margin checks on the first active section with non-excluded content
            if current_page not in excluded_set and not margin_checked and doc.sections:
                sec = doc.sections[current_sec_idx]
                actual_cols = get_column_count(sec)
                if actual_cols != target_cols:
                    errors.append(f"Column layout mismatch: Expected {target_cols}-column layout, detected {actual_cols}-column.")
                else:
                    passed.append(f"Layout geometry matches {target_cols}-column specification.")

                actual_margins = {
                    "top": round(sec.top_margin.inches, 2) if sec.top_margin else 0,
                    "bottom": round(sec.bottom_margin.inches, 2) if sec.bottom_margin else 0,
                    "left": round(sec.left_margin.inches, 2) if sec.left_margin else 0,
                    "right": round(sec.right_margin.inches, 2) if sec.right_margin else 0,
                }
                for side, expected in target_margins.items():
                    actual = actual_margins.get(side, 0)
                    if abs(actual - expected) > 0.05:
                        errors.append(f"{side.capitalize()} margin error: Target is {expected}\", document has {actual}\")")
                    else:
                        passed.append(f"{side.capitalize()} margin compliant ({expected}\")")
                margin_checked = True

            # Skip paragraph checking if current page is in excluded set
            if current_page in excluded_set:
                if page_breaks:
                    current_page += len(page_breaks)
                    paragraph_counter = 1
                    current_page_line_count = 0
                if has_section_break:
                    current_sec_idx = min(current_sec_idx + 1, len(doc.sections) - 1)
                continue

            valid_paragraphs += 1
            
            # Format text snippet for location context
            snippet = f"{text[:25]}..." if len(text) > 25 else text
            location_tag = f"Page {current_page}, Paragraph {paragraph_counter} (starting with '{snippet}')"

            chars_per_line = 75 if target_cols == 1 else 35
            estimated_lines = max(1, int(len(text) / chars_per_line) + 1)
            current_page_line_count += estimated_lines

            # Paragraph Classification Logic
            style_name = (p.style.name or "").lower() if p.style else ""
            is_bold_para = any(r.bold is True for r in p.runs if r.text.strip()) or (p.style and p.style.font and p.style.font.bold)
            is_italic_para = any(r.italic is True for r in p.runs if r.text.strip()) or (p.style and p.style.font and p.style.font.italic)
            
            run_sizes = [r.font.size.pt for r in p.runs if r.font and r.font.size and r.text.strip()]

            style_size = p.style.font.size.pt if (p.style and p.style.font and p.style.font.size) else None
            effective_sizes = run_sizes if run_sizes else ([style_size] if style_size is not None else [])

            text_clean = text.lower().strip()
            
            # Key Structural Patterns
            is_chapter_line = bool(re.match(r'^chapter\s+([0-9]+|[ivxdlcms]+|\w+)', text_clean))
            
            # 3-level section numbers
            is_numbered_subsubheading = bool(
                re.match(r'^\d+(?:\.\d+){2,}\.?\s*\w+|^[a-zA-Za-z]\.|\(?[a-zA-Za-z0-9]+\)\s*\w+', text_clean)
            )

            # Strictly 2-level section numbers
            is_numbered_subheading = bool(
                re.match(r'^\d+\.\d+\.?(?!\d|\.)\s*\w+', text_clean)
            )

            is_reference_line = bool(re.match(r'^\[\d+(?:--?\d+)?\]', text_clean))
            
            is_all_caps = text.isupper() and len(text) <= 65 and not text.endswith(('.', ';', ':'))
            is_short_bold_line = len(text) <= 80 and is_bold_para and not text.endswith(('.', ';', ':', ','))

            # Captions Classification
            caption_match = re.match(r'^(figure|fig\.|table)\s*[\d\.\w]+(.*)', text_clean)
            is_caption = False
            if caption_match:
                remainder = caption_match.group(2).strip()
                is_in_text_reference = remainder.startswith(BODY_REFERENCE_VERBS)
                
                if not is_in_text_reference:
                    if remainder == "" or remainder.startswith((':', '-', '.', '—', '–')) or len(text) <= 80:
                        is_caption = True

            # 1. Main Title / Chapter Heading Classification
            is_title = not is_caption and not is_reference_line and not is_numbered_subheading and not is_numbered_subsubheading and (
                "title" in style_name or
                style_name == "heading 1" or
                is_chapter_line or
                (prev_was_chapter and (is_all_caps or is_bold_para or len(text) <= 65)) or
                (is_all_caps and (p.alignment == WD_ALIGN_PARAGRAPH.CENTER or is_bold_para)) or
                (current_page == 1 and paragraph_counter == 1 and any(abs(sz - target_title_size) < 1.0 for sz in effective_sizes))
            )

            # 2. Level 3 Sub-subheading Classification
            is_subsubheading = not is_caption and not is_reference_line and not is_title and (
                "heading 3" in style_name or
                "heading 4" in style_name or
                is_numbered_subsubheading or
                (
                    any(abs(sz - target_subsub_size) < 1.0 for sz in effective_sizes) and
                    (not require_subsub_bold or is_bold_para) and
                    (not require_subsub_italic or is_italic_para) and
                    not is_numbered_subheading
                )
            )

            # 3. Level 2 Subheading Classification
            is_subheading = not is_caption and not is_reference_line and not is_title and not is_subsubheading and (
                "heading 2" in style_name or
                "subtitle" in style_name or
                is_numbered_subheading or
                text_clean in KNOWN_HEADINGS or
                is_short_bold_line
            )

            prev_was_chapter = is_chapter_line

            if is_caption:
                has_captions = True
                align = p.alignment
                if align is None and p.style and p.style.paragraph_format:
                    align = p.style.paragraph_format.alignment

                if align != WD_ALIGN_PARAGRAPH.CENTER:
                    errors.append(f"{location_tag} [Caption]: Caption '{snippet}' must be center-aligned.")
                    caption_clean = False

            elif is_title:
                for run in p.runs:
                    if run.font.name and run.font.name.lower() != target_title_font.lower():
                        errors.append(f"{location_tag} [Title]: Font '{run.font.name}' does not match target title font '{target_title_font}'.")
                        title_clean = False
                        break
                    if run.font.size and abs(run.font.size.pt - target_title_size) > 0.5:
                        errors.append(f"{location_tag} [Title]: Font size {run.font.size.pt}pt does not match target title size {target_title_size}pt.")
                        title_clean = False
                        break

            elif is_subsubheading:
                has_subsubs = True
                if require_subsub_bold and not is_bold_para:
                    errors.append(f"{location_tag} [Sub-subheading]: Level 3 subheadings are required to be Bold.")
                    subsub_clean = False
                if require_subsub_italic and not is_italic_para:
                    errors.append(f"{location_tag} [Sub-subheading]: Level 3 subheadings are required to be Italic.")
                    subsub_clean = False

                for run in p.runs:
                    if run.font.name and run.font.name.lower() != target_subsub_font.lower():
                        errors.append(f"{location_tag} [Sub-subheading]: Font '{run.font.name}' does not match target font '{target_subsub_font}'.")
                        subsub_clean = False
                        break
                    if run.font.size and abs(run.font.size.pt - target_subsub_size) > 0.5:
                        errors.append(f"{location_tag} [Sub-subheading]: Font size {run.font.size.pt}pt does not match target size {target_subsub_size}pt.")
                        subsub_clean = False
                        break

            elif is_subheading:
                if require_subtitle_bold and not is_bold_para:
                    errors.append(f"{location_tag} [Subheading]: Subheadings are required to be Bold.")
                    subtitle_clean = False

                for run in p.runs:
                    if run.font.name and run.font.name.lower() != target_subtitle_font.lower():
                        errors.append(f"{location_tag} [Subheading]: Font '{run.font.name}' does not match target heading font '{target_subtitle_font}'.")
                        subtitle_clean = False
                        break
                    if run.font.size and abs(run.font.size.pt - target_subtitle_size) > 0.5:
                        errors.append(f"{location_tag} [Subheading]: Font size {run.font.size.pt}pt does not match target heading size {target_subtitle_size}pt.")
                        subtitle_clean = False
                        break

            else:
                line_spacing = p.paragraph_format.line_spacing
                if line_spacing is not None and isinstance(line_spacing, (float, int)):
                    if abs(line_spacing - target_spacing) > 0.15:
                        errors.append(f"{location_tag}: Line spacing is {line_spacing}x (expected {target_spacing}x).")
                        spacing_clean = False

                if preset == "apa7":
                    indent = p.paragraph_format.first_line_indent
                    if indent is not None and abs(indent.inches - 0.5) > 0.1:
                        errors.append(f"{location_tag}: Missing standard 0.5-inch first-line indent.")
                        indent_clean = False

                if is_true(check_text_justification):
                    if p.alignment is not None and p.alignment != target_word_align_enum:
                        errors.append(f"{location_tag}: Paragraph alignment does not match required {target_alignment} format.")
                        justification_clean = False

                for run in p.runs:
                    if run.font.name and run.font.name.lower() != target_font.lower():
                        errors.append(f"{location_tag}: Font '{run.font.name}' does not match target body font '{target_font}'.")
                        font_clean = False
                        break

                for run in p.runs:
                    if run.font.size and abs(run.font.size.pt - target_size) > 0.5:
                        errors.append(f"{location_tag}: Font size {run.font.size.pt}pt does not match target body size {target_size}pt.")
                        size_clean = False
                        break

            paragraph_counter += 1

            if page_breaks:
                current_page += len(page_breaks)
                paragraph_counter = 1
                current_page_line_count = 0
            elif current_page_line_count >= max_lines_per_page:
                current_page += 1
                paragraph_counter = 1
                current_page_line_count = estimated_lines

            if has_section_break:
                current_sec_idx = min(current_sec_idx + 1, len(doc.sections) - 1)

        if preset == "custom":
            if title_clean and valid_paragraphs > 0:
                passed.append(f"Document Title typography matches target ({target_title_font}, {target_title_size}pt).")
            if subtitle_clean and valid_paragraphs > 0:
                passed.append(f"Subheadings/Section Headers match target ({target_subtitle_font}, {target_subtitle_size}pt{' Bold' if require_subtitle_bold else ''}).")
            if has_subsubs and subsub_clean:
                styles = []
                if require_subsub_bold:
                    styles.append("Bold")
                if require_subsub_italic:
                    styles.append("Italic")
                style_str = f" ({', '.join(styles)})" if styles else ""
                passed.append(f"Sub-subheadings match target ({target_subsub_font}, {target_subsub_size}pt{style_str}).")

        if has_captions and caption_clean:
            passed.append("Table and Figure captions comply with required center alignment.")
        if spacing_clean and valid_paragraphs > 0:
            passed.append(f"Line spacing matches target standard ({target_spacing}x).")
        if preset == "apa7" and indent_clean and valid_paragraphs > 0:
            passed.append("First-line indents comply with APA 0.5-inch requirement.")
        if is_true(check_text_justification) and justification_clean and valid_paragraphs > 0:
            passed.append(f"Paragraph alignment complies with required setting ({target_alignment}).")
        if font_clean and valid_paragraphs > 0:
            passed.append(f"Typography matches target font family ({target_font}).")
        if size_clean and valid_paragraphs > 0:
            passed.append(f"Font size matches target standard ({target_size} pt).")

        total_checks = len(errors) + len(passed)
        score = 100 if total_checks == 0 else int((len(passed) / total_checks) * 100)

        return {
            "preset_name": preset_label,
            "compliance_score": score,
            "errors": errors,
            "passed": passed
        }

    else:
        raise HTTPException(
            status_code=400, 
            detail="Unsupported file format. Please upload a .docx, .ppt, or .pptx file."
        )